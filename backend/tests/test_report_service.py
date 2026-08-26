import io
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from automation.organizer import process_file
from services.report_service import build_csv, build_excel, build_pdf


def _pdf_text(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    return "\n".join(page.extract_text() for page in reader.pages)


def _seed(watch_folder: Path, db_session, test_user):
    (watch_folder / "notes.txt").write_text("hi")
    process_file(watch_folder / "notes.txt", db_session, test_user)
    (watch_folder / "script.py").write_text("print(1)")
    process_file(watch_folder / "script.py", db_session, test_user)


def test_build_csv_contains_header_and_rows(watch_folder, test_rules, db_session, test_user):
    _seed(watch_folder, db_session, test_user)

    content = build_csv(db_session, test_user["id"]).decode("utf-8")

    assert "Timestamp" in content
    assert "notes.txt" in content
    assert "script.py" in content


def test_build_excel_produces_nonempty_workbook(watch_folder, test_rules, db_session, test_user):
    _seed(watch_folder, db_session, test_user)

    content = build_excel(db_session, test_user["id"])

    # .xlsx files are zip archives; a real workbook is well over a trivial size
    assert content[:2] == b"PK"
    assert len(content) > 1000


def test_build_pdf_produces_a_pdf_document(watch_folder, test_rules, db_session, test_user):
    _seed(watch_folder, db_session, test_user)

    content = build_pdf(db_session, test_user["id"])

    assert content.startswith(b"%PDF")


def test_build_pdf_actually_contains_the_title_categories_and_totals(
    watch_folder, test_rules, db_session, test_user
):
    """A successful PDF byte response isn't proof the content is correct --
    this decodes the actual page text (via pypdf) rather than trusting the
    %PDF magic-byte check alone."""
    _seed(watch_folder, db_session, test_user)

    text = _pdf_text(build_pdf(db_session, test_user["id"]))

    assert "FileFlow Report" in text
    assert "Total files organized" in text
    # notes.txt -> Documents, script.py -> Python Projects, per test_rules'
    # category mapping -- both category names must appear in the "By
    # category" breakdown.
    assert "Documents" in text
    assert "Python Projects" in text


def test_reports_do_not_error_on_empty_database(db_session):
    assert build_csv(db_session, 1)
    assert build_excel(db_session, 1)
    assert build_pdf(db_session, 1)


class TestFormulaInjection:
    """A file's name comes entirely from what the user (or whoever placed
    the file in their watched folder) named it on disk -- report_service.py
    writes original_filename/final_filename/category straight into CSV and
    Excel cells with no escaping. A name starting with =, +, -, or @ is a
    live formula to Excel/Sheets/LibreOffice, not literal text (CWE-1236).
    For .xlsx specifically this isn't just a "if a human retypes it" risk:
    openpyxl itself inspects any string starting with "=" and marks the
    cell's data_type as a formula ('f') when writing the file, so the
    injected formula is already a real formula cell before the file is
    ever opened."""

    def _log(self, db_session, user_id, name, category="Documents"):
        db_session.activity_log.insert_one(
            {
                "user_id": user_id,
                "timestamp": datetime.now(timezone.utc),
                "original_filename": name,
                "final_filename": name,
                "category": category,
                "action": "moved",
                "file_size": 10,
            }
        )

    def test_csv_neutralizes_a_formula_prefixed_filename(self, db_session):
        self._log(db_session, 1, "=cmd|' /C calc'!A0.txt")

        content = build_csv(db_session, 1).decode("utf-8")

        # The raw formula must never appear as a bare, un-neutralized field --
        # a spreadsheet program must see this as text, not something to
        # evaluate.
        assert ",=cmd|' /C calc'!A0.txt," not in content
        assert "'=cmd|' /C calc'!A0.txt" in content

    def test_excel_cell_is_not_flagged_as_a_formula(self, db_session):
        self._log(db_session, 1, "=1+1")

        content = build_excel(db_session, 1)
        wb = load_workbook(io.BytesIO(content))
        cell = wb["Activity Log"]["B2"]

        # openpyxl sets data_type == "f" for any string it writes that
        # starts with "=" -- confirming data_type is "s" (a plain string)
        # proves the sanitized value no longer reads as a formula to
        # openpyxl itself, and therefore to Excel when it opens the file.
        assert cell.data_type != "f"
        assert cell.value == "'=1+1"

    def test_plus_minus_and_at_prefixes_are_also_neutralized_in_csv(self, db_session):
        self._log(db_session, 1, "+1+1", category="Documents")
        self._log(db_session, 1, "-1+1", category="Documents")
        self._log(db_session, 1, "@SUM(A1)", category="Documents")

        content = build_csv(db_session, 1).decode("utf-8")

        assert "'+1+1" in content
        assert "'-1+1" in content
        assert "'@SUM(A1)" in content

    def test_a_malicious_category_name_is_also_neutralized(self, db_session):
        self._log(db_session, 1, "invoice.pdf", category="=HYPERLINK(\"http://evil\")")

        csv_content = build_csv(db_session, 1).decode("utf-8")
        xlsx = build_excel(db_session, 1)
        wb = load_workbook(io.BytesIO(xlsx))

        assert "'=HYPERLINK" in csv_content
        assert wb["Activity Log"]["D2"].data_type != "f"

    def test_an_ordinary_filename_is_left_completely_unchanged(self, db_session):
        self._log(db_session, 1, "notes.txt")

        content = build_csv(db_session, 1).decode("utf-8")

        assert "notes.txt" in content
        assert "'notes.txt" not in content


class TestUserIsolation:
    """_log_rows()/get_stats() both scope every query by user_id -- these
    confirm that scoping actually holds for the three export formats
    themselves, not just for the /api/logs and /api/stats routes already
    covered elsewhere."""

    def test_csv_never_contains_another_users_filenames(
        self, watch_folder, test_rules, db_session, test_user
    ):
        _seed(watch_folder, db_session, test_user)
        db_session.activity_log.insert_one(
            {
                "user_id": 999,
                "timestamp": datetime.now(timezone.utc),
                "original_filename": "someone_elses_secret_file.txt",
                "final_filename": "someone_elses_secret_file.txt",
                "category": "Documents",
                "action": "moved",
                "file_size": 5,
            }
        )

        content = build_csv(db_session, test_user["id"]).decode("utf-8")

        assert "notes.txt" in content
        assert "someone_elses_secret_file.txt" not in content

    def test_excel_never_contains_another_users_filenames_or_stats(
        self, watch_folder, test_rules, db_session, test_user
    ):
        _seed(watch_folder, db_session, test_user)
        db_session.activity_log.insert_one(
            {
                "user_id": 999,
                "timestamp": datetime.now(timezone.utc),
                "original_filename": "someone_elses_secret_file.txt",
                "final_filename": "someone_elses_secret_file.txt",
                "category": "Documents",
                "action": "moved",
                "file_size": 5,
            }
        )

        content = build_excel(db_session, test_user["id"])
        wb = load_workbook(io.BytesIO(content))
        log_sheet_values = [
            cell.value for row in wb["Activity Log"].iter_rows() for cell in row
        ]

        assert "someone_elses_secret_file.txt" not in log_sheet_values
        # the summary "Total files organized" row must count only this
        # account's own two seeded files, not the other account's one.
        summary_values = {
            row[0].value: row[1].value
            for row in wb["Summary"].iter_rows()
            if row[0].value == "Total files organized"
        }
        assert summary_values["Total files organized"] == 2

    def test_pdf_never_contains_another_users_filename(
        self, watch_folder, test_rules, db_session, test_user
    ):
        _seed(watch_folder, db_session, test_user)
        db_session.activity_log.insert_one(
            {
                "user_id": 999,
                "timestamp": datetime.now(timezone.utc),
                "original_filename": "someone_elses_secret_file.txt",
                "final_filename": "someone_elses_secret_file.txt",
                "category": "Documents",
                "action": "moved",
                "file_size": 5,
            }
        )

        content = build_pdf(db_session, test_user["id"])

        assert "someone_elses_secret_file.txt" not in _pdf_text(content)


class TestMultiWatchFolderAggregation:
    """Reporting is account-scoped, not folder-scoped (established in the
    prior multi-watch-folder audit) -- confirms activity from an additional
    watch folder is actually aggregated into this account's reports, not
    silently dropped because it happened outside the primary folder."""

    def test_csv_and_stats_include_activity_from_an_additional_folder(
        self, watch_folder, test_rules, db_session, test_user, tmp_path
    ):
        from automation.organizer import organize_existing_files

        extra_folder = tmp_path / "extra_watch"
        extra_folder.mkdir()
        user = {**test_user, "additional_watch_folders": [str(extra_folder)]}

        (watch_folder / "primary_file.txt").write_text("hi")
        (extra_folder / "extra_file.txt").write_text("hi")

        result = organize_existing_files(db_session, user)
        assert result["organized"] == 2

        csv_content = build_csv(db_session, user["id"]).decode("utf-8")
        assert "primary_file.txt" in csv_content
        assert "extra_file.txt" in csv_content

        xlsx = build_excel(db_session, user["id"])
        wb = load_workbook(io.BytesIO(xlsx))
        summary_values = {
            row[0].value: row[1].value
            for row in wb["Summary"].iter_rows()
            if row[0].value == "Total files organized"
        }
        assert summary_values["Total files organized"] == 2


class TestCrossFormatConsistency:
    """The main audit objective: given the same underlying activity, CSV,
    Excel, and PDF must agree on the same totals -- they all read through
    the same _log_rows()/get_stats() functions, so this pins that down
    with an actual reproduction rather than assuming it from the shared
    code path alone."""

    def test_csv_excel_and_pdf_report_the_same_file_count_and_names(
        self, watch_folder, test_rules, db_session, test_user
    ):
        _seed(watch_folder, db_session, test_user)

        csv_content = build_csv(db_session, test_user["id"]).decode("utf-8")
        xlsx = build_excel(db_session, test_user["id"])
        pdf_text = _pdf_text(build_pdf(db_session, test_user["id"]))
        wb = load_workbook(io.BytesIO(xlsx))

        for name in ("notes.txt", "script.py"):
            assert name in csv_content
            assert name in pdf_text
            log_values = [
                cell.value for row in wb["Activity Log"].iter_rows() for cell in row
            ]
            assert name in log_values

        summary_values = {
            row[0].value: row[1].value
            for row in wb["Summary"].iter_rows()
            if row[0].value == "Total files organized"
        }
        assert summary_values["Total files organized"] == 2
