import csv
import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from services.stats_service import get_stats

LOG_HEADERS = [
    "Timestamp",
    "Original Filename",
    "Final Filename",
    "Category",
    "Action",
    "Size (bytes)",
]


def _log_rows(
    db,
    user_id: int,
    limit: int = 500,
    since: datetime | None = None,
    until: datetime | None = None,
) -> list[list]:
    query: dict = {"user_id": user_id}
    if since is not None or until is not None:
        query["timestamp"] = {
            **({"$gte": since} if since is not None else {}),
            **({"$lt": until} if until is not None else {}),
        }
    logs = db.activity_log.find(query).sort("_id", -1).limit(limit)
    return [
        [
            log["timestamp"].strftime("%Y-%m-%d %H:%M:%S"),
            log["original_filename"],
            log["final_filename"],
            log["category"],
            log["action"],
            log["file_size"],
        ]
        for log in logs
    ]


def _generated_at() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")


def build_csv(
    db, user_id: int, since: datetime | None = None, until: datetime | None = None
) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(LOG_HEADERS)
    writer.writerows(_log_rows(db, user_id, since=since, until=until))
    return buffer.getvalue().encode("utf-8")


def build_excel(db, user_id: int) -> bytes:
    stats = get_stats(db, user_id)
    bold = Font(bold=True)

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"

    summary.append(["FileFlow Report"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([f"Generated {_generated_at()}"])
    summary.append([])

    header_row = summary.max_row + 1
    summary.append(["Metric", "Value"])
    for cell in summary[header_row]:
        cell.font = bold

    for label, value in [
        ("Total files organized", stats.total_files_organized),
        ("Duplicates detected", stats.duplicates_detected),
        ("Files archived", stats.files_archived),
        ("Storage saved (bytes)", stats.storage_saved_bytes),
        ("Organized today", stats.today_count),
    ]:
        summary.append([label, value])

    summary.append([])
    header_row = summary.max_row + 1
    summary.append(["Category", "Count"])
    for cell in summary[header_row]:
        cell.font = bold
    for cat in stats.by_category:
        summary.append([cat.category, cat.count])

    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 16

    log_sheet = wb.create_sheet("Activity Log")
    log_sheet.append(LOG_HEADERS)
    for cell in log_sheet[1]:
        cell.font = bold
    for row in _log_rows(db, user_id):
        log_sheet.append(row)
    for i, width in enumerate([20, 28, 28, 16, 20, 14], start=1):
        log_sheet.column_dimensions[get_column_letter(i)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _styled_table(data, col_widths=None):
    table = Table(data, colWidths=col_widths)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2a78d6")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e1e0d9")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f9f9f7")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def build_pdf(db, user_id: int) -> bytes:
    stats = get_stats(db, user_id)
    styles = getSampleStyleSheet()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=letter, topMargin=0.6 * inch, bottomMargin=0.6 * inch
    )

    story = [
        Paragraph("FileFlow Report", styles["Title"]),
        Paragraph(f"Generated {_generated_at()}", styles["Normal"]),
        Spacer(1, 0.25 * inch),
    ]

    summary_data = [
        ["Metric", "Value"],
        ["Total files organized", stats.total_files_organized],
        ["Duplicates detected", stats.duplicates_detected],
        ["Files archived", stats.files_archived],
        ["Storage saved (bytes)", stats.storage_saved_bytes],
        ["Organized today", stats.today_count],
    ]
    story.append(_styled_table(summary_data))
    story.append(Spacer(1, 0.25 * inch))

    if stats.by_category:
        story.append(Paragraph("By category", styles["Heading2"]))
        category_data = [["Category", "Count"]] + [
            [c.category, c.count] for c in stats.by_category
        ]
        story.append(_styled_table(category_data))
        story.append(Spacer(1, 0.25 * inch))

    story.append(Paragraph("Recent activity", styles["Heading2"]))
    log_rows = _log_rows(db, user_id, limit=50)
    if log_rows:
        table_data = [LOG_HEADERS] + log_rows
        story.append(_styled_table(table_data, col_widths=[95, 95, 95, 65, 75, 60]))
    else:
        story.append(Paragraph("No activity recorded yet.", styles["Normal"]))

    doc.build(story)
    return buffer.getvalue()
